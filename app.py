"""AI Resume ATS Checker - Streamlit + Google Gemini Flash.

Upload a resume (PDF or DOCX), optionally paste a job description, and get:
  * an ATS score (0-100) with a category breakdown
  * strengths, problems, missing keywords
  * concrete, prioritized improvements and rewrite suggestions
"""

from __future__ import annotations

import io
import json
import os
import re
from typing import Any

import streamlit as st

# --------------------------------------------------------------------------- #
# Configuration
# --------------------------------------------------------------------------- #

# Tried in order. The user's choice goes first; the rest act as fallbacks if a
# model ID is unavailable for the API key / region.
MODEL_OPTIONS = [
    "gemini-3.5-flash",
    "gemini-3-flash-preview",
    "gemini-2.5-flash",
]

MAX_FILE_MB = 5
MAX_RESUME_CHARS = 15000
MAX_JD_CHARS = 6000
MIN_RESUME_CHARS = 200

# Category -> weight (sums to 100). The overall score is computed here in code
# from the category scores so it is consistent and explainable.
CATEGORY_WEIGHTS = {
    "keyword_match": 30,
    "content_quality": 25,
    "formatting_parsability": 20,
    "structure_sections": 15,
    "readability_grammar": 10,
}

CATEGORY_LABELS = {
    "keyword_match": "Keywords & relevance",
    "content_quality": "Content quality & impact",
    "formatting_parsability": "ATS formatting / parsability",
    "structure_sections": "Structure & sections",
    "readability_grammar": "Readability & grammar",
}

SYSTEM_INSTRUCTION = """You are an expert ATS (Applicant Tracking System) analyst and \
professional resume reviewer. You evaluate resumes strictly and honestly, never \
inflating scores. You only use information present in the resume text; never invent \
experience, employers, or numbers. Always answer with valid JSON only."""

PROMPT_TEMPLATE = """Analyze the resume below for ATS compatibility and quality.

{jd_block}

Score each category from 0 to 100 (integers):
- keyword_match: relevant skills/keywords {keyword_hint}
- content_quality: achievements, quantified impact, action verbs, clarity
- formatting_parsability: clean text extraction, standard headings, dates, contact info, \
no signs of tables/columns/graphics breaking parsing
- structure_sections: presence/order of Contact, Summary, Experience, Education, Skills, \
Projects/Certifications as appropriate
- readability_grammar: grammar, tense consistency, concision, length

Return ONLY a JSON object with exactly this shape:
{{
  "candidate_role_guess": "string",
  "scores": {{
    "keyword_match": 0,
    "content_quality": 0,
    "formatting_parsability": 0,
    "structure_sections": 0,
    "readability_grammar": 0
  }},
  "summary": "2-3 sentence overall assessment",
  "strengths": ["..."],
  "issues": [
    {{"severity": "high|medium|low", "problem": "...", "fix": "..."}}
  ],
  "missing_keywords": ["..."],
  "present_keywords": ["..."],
  "rewrite_suggestions": [
    {{"original": "text copied from the resume", "improved": "stronger rewrite without inventing facts", "why": "..."}}
  ],
  "action_plan": ["prioritized, concrete step 1", "step 2", "step 3"]
}}

Rules: 3-6 strengths, 4-8 issues sorted by severity, up to 12 missing_keywords, \
up to 12 present_keywords, 3-5 rewrite_suggestions, 4-7 action_plan steps. \
In rewrite_suggestions, if a metric is unknown use a placeholder like [X%] instead of \
making up numbers.

RESUME TEXT:
\"\"\"
{resume}
\"\"\"
"""

# --------------------------------------------------------------------------- #
# File parsing
# --------------------------------------------------------------------------- #


class ResumeError(Exception):
    """User-facing error (bad file, empty text, API problem, ...)."""


def extract_text_from_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception:
                raise ResumeError("This PDF is password-protected. Please remove the password and try again.")
        pages = [(page.extract_text() or "") for page in reader.pages]
    except ResumeError:
        raise
    except Exception as exc:  # corrupt file etc.
        raise ResumeError(f"Could not read the PDF: {exc}")
    return "\n".join(pages)


def extract_text_from_docx(data: bytes) -> str:
    from docx import Document

    try:
        doc = Document(io.BytesIO(data))
    except Exception as exc:
        raise ResumeError(f"Could not read the DOCX file: {exc}")

    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    # Many resumes keep content in tables; include them.
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def clean_text(text: str) -> str:
    text = text.replace("\x00", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def extract_resume_text(filename: str, data: bytes) -> str:
    name = filename.lower()
    if len(data) > MAX_FILE_MB * 1024 * 1024:
        raise ResumeError(f"File is larger than {MAX_FILE_MB} MB.")
    if name.endswith(".pdf"):
        text = extract_text_from_pdf(data)
    elif name.endswith(".docx"):
        text = extract_text_from_docx(data)
    elif name.endswith(".txt"):
        text = data.decode("utf-8", errors="ignore")
    else:
        raise ResumeError("Unsupported file type. Upload a PDF, DOCX or TXT file.")

    text = clean_text(text)
    if len(text) < MIN_RESUME_CHARS:
        raise ResumeError(
            "Almost no text could be extracted. If your resume is a scanned image or "
            "heavily graphic, an ATS cannot read it either - export a text-based PDF/DOCX."
        )
    return text


# --------------------------------------------------------------------------- #
# Gemini call + result handling
# --------------------------------------------------------------------------- #


def build_prompt(resume_text: str, job_description: str = "") -> str:
    resume_text = resume_text[:MAX_RESUME_CHARS]
    jd = job_description.strip()[:MAX_JD_CHARS]
    if jd:
        jd_block = (
            "A target JOB DESCRIPTION is provided. Judge keyword_match against it and "
            "list missing keywords from it.\n\nJOB DESCRIPTION:\n\"\"\"\n" + jd + "\n\"\"\""
        )
        keyword_hint = "that match the job description (hard skills, tools, titles, certifications)"
    else:
        jd_block = "No job description was provided; judge against general standards for the candidate's apparent target role."
        keyword_hint = "typical for the candidate's apparent target role (hard skills, tools, titles)"
    return PROMPT_TEMPLATE.format(jd_block=jd_block, keyword_hint=keyword_hint, resume=resume_text)


def parse_json_response(raw: str) -> dict[str, Any]:
    """Parse model output into a dict, tolerating ```json fences and stray text."""
    if not raw or not raw.strip():
        raise ResumeError("The model returned an empty response. Please try again.")
    text = raw.strip()
    text = re.sub(r"^```(?:json)?\s*", "", text)
    text = re.sub(r"\s*```$", "", text)
    try:
        data = json.loads(text)
    except json.JSONDecodeError:
        start, end = text.find("{"), text.rfind("}")
        if start == -1 or end <= start:
            raise ResumeError("The model response was not valid JSON. Please try again.")
        try:
            data = json.loads(text[start : end + 1])
        except json.JSONDecodeError:
            raise ResumeError("The model response was not valid JSON. Please try again.")
    if not isinstance(data, dict):
        raise ResumeError("Unexpected model response format. Please try again.")
    return data


def _clamp_score(value: Any) -> int:
    try:
        return max(0, min(100, int(round(float(value)))))
    except (TypeError, ValueError):
        return 0


def _str_list(value: Any, limit: int = 20) -> list[str]:
    if not isinstance(value, list):
        return []
    return [str(v).strip() for v in value if str(v).strip()][:limit]


def normalize_result(data: dict[str, Any]) -> dict[str, Any]:
    """Validate/clean model output and compute the weighted overall score."""
    raw_scores = data.get("scores") if isinstance(data.get("scores"), dict) else {}
    scores = {key: _clamp_score(raw_scores.get(key, 0)) for key in CATEGORY_WEIGHTS}
    overall = round(sum(scores[k] * w for k, w in CATEGORY_WEIGHTS.items()) / 100)

    issues = []
    for item in data.get("issues", []) if isinstance(data.get("issues"), list) else []:
        if not isinstance(item, dict):
            continue
        severity = str(item.get("severity", "medium")).lower()
        if severity not in ("high", "medium", "low"):
            severity = "medium"
        issues.append(
            {
                "severity": severity,
                "problem": str(item.get("problem", "")).strip(),
                "fix": str(item.get("fix", "")).strip(),
            }
        )
    order = {"high": 0, "medium": 1, "low": 2}
    issues = sorted([i for i in issues if i["problem"]], key=lambda i: order[i["severity"]])

    rewrites = []
    for item in data.get("rewrite_suggestions", []) if isinstance(data.get("rewrite_suggestions"), list) else []:
        if isinstance(item, dict) and item.get("improved"):
            rewrites.append(
                {
                    "original": str(item.get("original", "")).strip(),
                    "improved": str(item.get("improved", "")).strip(),
                    "why": str(item.get("why", "")).strip(),
                }
            )

    return {
        "overall": overall,
        "scores": scores,
        "role": str(data.get("candidate_role_guess", "")).strip(),
        "summary": str(data.get("summary", "")).strip(),
        "strengths": _str_list(data.get("strengths")),
        "issues": issues,
        "missing_keywords": _str_list(data.get("missing_keywords")),
        "present_keywords": _str_list(data.get("present_keywords")),
        "rewrites": rewrites,
        "action_plan": _str_list(data.get("action_plan")),
    }


def score_band(score: int) -> tuple[str, str]:
    if score >= 80:
        return "Excellent", "green"
    if score >= 65:
        return "Good", "blue"
    if score >= 50:
        return "Needs work", "orange"
    return "Poor", "red"


def call_gemini(api_key: str, models: list[str], prompt: str) -> tuple[str, str]:
    """Call Gemini, falling back through `models`. Returns (raw_text, model_used)."""
    from google import genai
    from google.genai import errors, types

    client = genai.Client(api_key=api_key)
    config = types.GenerateContentConfig(
        system_instruction=SYSTEM_INSTRUCTION,
        temperature=0.2,
        response_mime_type="application/json",
    )
    last_error: Exception | None = None
    for model in models:
        try:
            response = client.models.generate_content(model=model, contents=prompt, config=config)
            text = getattr(response, "text", None)
            if text:
                return text, model
            last_error = ResumeError(f"Model {model} returned no text (possibly blocked).")
        except errors.APIError as exc:
            code = getattr(exc, "code", None)
            if code in (401, 403) or "API key" in str(exc):
                raise ResumeError("The Gemini API key was rejected. Check that it is valid.")
            last_error = exc  # 404 / 429 / 5xx -> try the next model
        except Exception as exc:
            last_error = exc
    raise ResumeError(f"Gemini request failed for all models. Last error: {last_error}")


def analyze_resume(api_key: str, models: list[str], resume_text: str, job_description: str = "") -> dict[str, Any]:
    raw, model_used = call_gemini(api_key, models, build_prompt(resume_text, job_description))
    result = normalize_result(parse_json_response(raw))
    result["model_used"] = model_used
    return result


def report_markdown(result: dict[str, Any]) -> str:
    label, _ = score_band(result["overall"])
    lines = [f"# ATS Resume Report", "", f"**Overall ATS score: {result['overall']}/100 ({label})**", ""]
    if result["summary"]:
        lines += [result["summary"], ""]
    lines += ["## Category scores"]
    lines += [f"- {CATEGORY_LABELS[k]}: {v}/100" for k, v in result["scores"].items()]
    lines += ["", "## Strengths"] + [f"- {s}" for s in result["strengths"]]
    lines += ["", "## Issues & fixes"]
    lines += [f"- [{i['severity'].upper()}] {i['problem']} -> {i['fix']}" for i in result["issues"]]
    lines += ["", "## Missing keywords", ", ".join(result["missing_keywords"]) or "None found"]
    lines += ["", "## Rewrite suggestions"]
    for r in result["rewrites"]:
        lines += [f"- Original: {r['original']}", f"  Improved: {r['improved']}", f"  Why: {r['why']}"]
    lines += ["", "## Action plan"] + [f"{n}. {s}" for n, s in enumerate(result["action_plan"], 1)]
    return "\n".join(lines) + "\n"


# --------------------------------------------------------------------------- #
# UI
# --------------------------------------------------------------------------- #


def get_api_key() -> str:
    try:
        key = st.secrets.get("GEMINI_API_KEY", "")
    except Exception:
        key = ""
    return key or os.environ.get("GEMINI_API_KEY", "")


def render_results(result: dict[str, Any]) -> None:
    label, color = score_band(result["overall"])
    top_left, top_right = st.columns([1, 2])
    with top_left:
        st.metric("ATS score", f"{result['overall']} / 100")
        st.markdown(f":{color}[**{label}**]")
        if result["role"]:
            st.caption(f"Detected target role: {result['role']}")
    with top_right:
        st.write(result["summary"])

    st.subheader("Score breakdown")
    for key, value in result["scores"].items():
        st.write(f"{CATEGORY_LABELS[key]} (weight {CATEGORY_WEIGHTS[key]}%) - **{value}**")
        st.progress(value / 100)

    tab_fix, tab_kw, tab_rewrite, tab_plan = st.tabs(
        ["Issues & fixes", "Keywords", "Rewrite suggestions", "Action plan"]
    )
    with tab_fix:
        if result["strengths"]:
            st.markdown("**Strengths**")
            for s in result["strengths"]:
                st.success(s)
        st.markdown("**Improvements**")
        icons = {"high": "🔴", "medium": "🟠", "low": "🟡"}
        for item in result["issues"]:
            with st.expander(f"{icons[item['severity']]} {item['problem']}"):
                st.write(f"**Fix:** {item['fix']}")
        if not result["issues"]:
            st.info("No issues reported.")
    with tab_kw:
        c1, c2 = st.columns(2)
        with c1:
            st.markdown("**Found in resume**")
            st.write(", ".join(result["present_keywords"]) or "None listed")
        with c2:
            st.markdown("**Missing / worth adding**")
            st.write(", ".join(result["missing_keywords"]) or "None listed")
    with tab_rewrite:
        if not result["rewrites"]:
            st.info("No rewrite suggestions returned.")
        for r in result["rewrites"]:
            st.markdown(f"**Original:** {r['original']}")
            st.markdown(f"**Improved:** {r['improved']}")
            if r["why"]:
                st.caption(r["why"])
            st.divider()
    with tab_plan:
        for n, step in enumerate(result["action_plan"], 1):
            st.markdown(f"{n}. {step}")

    st.download_button(
        "Download report (.md)",
        report_markdown(result),
        file_name="ats_resume_report.md",
        mime="text/markdown",
    )
    st.caption(f"Model used: {result.get('model_used', 'n/a')}")


def main() -> None:
    st.set_page_config(page_title="AI Resume ATS Checker", page_icon="📄", layout="wide")
    st.title("📄 AI Resume ATS Checker")
    st.write("Upload your resume to get an ATS score and specific ways to improve it.")

    with st.sidebar:
        st.header("Settings")
        api_key = get_api_key()
        if api_key:
            st.success("Gemini API key loaded.")
        else:
            api_key = st.text_input("Gemini API key", type="password", help="Get one free at aistudio.google.com")
        model = st.selectbox("Gemini model", MODEL_OPTIONS, help="Other models are used as fallbacks.")
        st.caption("Your resume is sent to the Gemini API for analysis and is not stored by this app.")

    col_a, col_b = st.columns(2)
    with col_a:
        uploaded = st.file_uploader("Resume (PDF, DOCX or TXT)", type=["pdf", "docx", "txt"])
    with col_b:
        job_description = st.text_area(
            "Target job description (optional, recommended)",
            height=160,
            placeholder="Paste the job posting here for a tailored keyword match...",
        )

    if st.button("Analyze resume", type="primary", disabled=uploaded is None):
        if not api_key:
            st.error("Please provide a Gemini API key in the sidebar.")
            st.stop()
        try:
            with st.spinner("Reading your resume..."):
                text = extract_resume_text(uploaded.name, uploaded.getvalue())
            models = [model] + [m for m in MODEL_OPTIONS if m != model]
            with st.spinner("Analyzing with Gemini..."):
                st.session_state["result"] = analyze_resume(api_key, models, text, job_description)
        except ResumeError as exc:
            st.session_state.pop("result", None)
            st.error(str(exc))

    if "result" in st.session_state:
        st.divider()
        render_results(st.session_state["result"])


if __name__ == "__main__":
    main()
